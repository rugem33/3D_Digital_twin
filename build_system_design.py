"""
양식 5 - 시스템 구조 설계 및 개발 환경 문서 생성 스크립트
템플릿의 스타일을 재사용하여 작성합니다.
"""
import shutil
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

TEMPLATE = "[양식 5] 시스템 구조 설계 및 개발 환경 양식.docx"
OUTPUT   = "[양식 5] 시스템 구조 설계 및 개발 환경 완성.docx"

shutil.copy(TEMPLATE, OUTPUT)
doc = Document(OUTPUT)

# ── 기존 내용 모두 삭제 ──────────────────────────────────────────────────────
for p in list(doc.paragraphs):
    p._element.getparent().remove(p._element)
for t in list(doc.tables):
    t._element.getparent().remove(t._element)

# ── 헬퍼 ─────────────────────────────────────────────────────────────────────
def add_paragraph(text, style=None, bold=False, size=None, color=None, align=None):
    p = doc.add_paragraph()
    if style:
        try:
            p.style = doc.styles[style]
        except Exception:
            pass
    run = p.add_run(text)
    if bold:
        run.bold = True
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    if align:
        p.alignment = align
    return p

def add_heading(text, level=1):
    style_map = {1: 'Heading 1', 2: 'Heading 2', 3: 'Heading 3'}
    try:
        p = doc.add_heading(text, level=level)
    except Exception:
        p = doc.add_paragraph(text)
        p.runs[0].bold = True
    return p

def add_table(headers, rows, col_widths=None):
    table = doc.add_table(rows=1+len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # 헤더 행
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = h
        for run in cell.paragraphs[0].runs:
            run.bold = True
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        shading = OxmlElement('w:shd')
        shading.set(qn('w:val'), 'clear')
        shading.set(qn('w:color'), 'auto')
        shading.set(qn('w:fill'), 'D9E1F2')
        cell._tc.get_or_add_tcPr().append(shading)

    # 데이터 행
    for r_idx, row_data in enumerate(rows):
        row = table.rows[r_idx + 1]
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.text = val
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT

    if col_widths:
        for r in table.rows:
            for i, cell in enumerate(r.cells):
                cell.width = Cm(col_widths[i])
    return table

# ════════════════════════════════════════════════════════════════════════════
# 표지
# ════════════════════════════════════════════════════════════════════════════
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("Project Document\nSystem Architecture Design\n& Development Environments")
run.bold = True
run.font.size = Pt(20)

doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("Project Name")
run.bold = True
run.font.size = Pt(14)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("Cesium for Unity 기반 모바일 1인칭 3D 지도")
run.bold = True
run.font.size = Pt(13)

doc.add_paragraph()
for line in [
    "15 조",
    "202002581  황용하",
    "202300385  이혜린",
    "202102664  여현서",
    "202102669  유성",
    "",
    "지도교수: 김형기 교수님",
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(line)

doc.add_page_break()

# ════════════════════════════════════════════════════════════════════════════
# Document Revision History
# ════════════════════════════════════════════════════════════════════════════
add_heading("Document Revision History", level=1)
add_table(
    headers=["Rev#", "Date", "Affected Section", "Author"],
    rows=[
        ["1", "2026/04/10", "Product Backlog 작성", "공통"],
        ["2", "2026/04/25", "시스템 구조 설계 및 개발 환경 작성", "공통"],
    ],
    col_widths=[2, 3, 7, 3],
)
doc.add_paragraph()

# ════════════════════════════════════════════════════════════════════════════
# 0. Meta Information
# ════════════════════════════════════════════════════════════════════════════
add_heading("System Design Document – Cesium for Unity 기반 모바일 1인칭 3D 지도", level=1)
add_heading("0. Meta Information", level=2)
add_table(
    headers=["항목", "내용"],
    rows=[
        ["Project",  "Cesium for Unity 기반 모바일 1인칭 3D 지도"],
        ["Team",     "15조 (황용하, 이혜린, 여현서, 유성)"],
        ["Version",  "0.1.0"],
        ["Engine",   "Unity 2022.3 LTS  ·  Universal Render Pipeline (URP)"],
        ["Platform", "Android / iOS (모바일)"],
        ["Scope",    "US-01 ~ US-15 (전체 7 Epic, 15 User Story)"],
    ],
    col_widths=[4, 12],
)
doc.add_paragraph()

# ════════════════════════════════════════════════════════════════════════════
# 1. Project Overview
# ════════════════════════════════════════════════════════════════════════════
add_heading("1. Project Overview", level=1)

add_heading("1.1 Vision", level=2)
doc.add_paragraph(
    "도보 및 자전거 이용자는 기존 2D 지도나 드론 뷰 기반 3D 지도로는 지면 밀착형 공간 정보(노면 상태, 가로등, "
    "자전거 전용 도로 등)를 직관적으로 확인하기 어렵다. 본 프로젝트는 Cesium for Unity 위에 1인칭 보행자 시점을 "
    "구현하고, GPS 실시간 동기화·커스텀 시설물 배치·자전거 도로 시각화를 통해 실감형 길찾기 경험을 모바일로 제공한다."
)

add_heading("1.2 Scope", level=2)
add_table(
    headers=["구분", "항목"],
    rows=[
        ["In-Scope",
         "3D 지형 스트리밍(Cesium Ion) / 실시간 GPS 동기화 / 1인칭 카메라 제어 / "
         "커스텀 시설물(가로등·벤치) 자동 배치 / 자전거 도로 폴리라인 시각화 / "
         "터치 기반 UI / 경로 가이드 시각화 / 모바일 성능 최적화"],
        ["Out-of-Scope",
         "결제 시스템 / 서버 사이드 사용자 인증 / 실시간 교통 정보 연동 / "
         "다중 사용자 협업 기능 / PC/콘솔 빌드"],
    ],
    col_widths=[4, 12],
)
doc.add_paragraph()

add_heading("1.3 Success Metrics", level=2)
metrics = [
    ("GPS-엔진 좌표 정밀도",   "WGS84 → Unity 3D 공간 좌표 매핑 오차 실외 기준 3 m 이내"),
    ("위치 업데이트 지연율",   "카메라 위치 업데이트 지연 500 ms 이내, Lerp 보간으로 이질감 최소화"),
    ("카메라 Jitter 발생률",   "GPS↔터치 모드 전환 시 카메라 튐 현상 1 % 미만"),
    ("데이터 렌더링 부하",     "GPS 이동 및 3D Tiles 스트리밍 중 최저 30 FPS 유지"),
    ("시설물/도로 가시 거리",  "자전거 도로 및 가로등 50 m 이상 가시 거리 확보"),
]
add_table(
    headers=["지표", "목표"],
    rows=[[k, v] for k, v in metrics],
    col_widths=[5, 11],
)
doc.add_paragraph()

# ════════════════════════════════════════════════════════════════════════════
# 2. Architecture Design
# ════════════════════════════════════════════════════════════════════════════
add_heading("2. Architecture Design", level=1)

add_heading("2.1 System Context Diagram", level=2)
doc.add_paragraph(
    "아래 표는 시스템의 주요 외부 Actor 및 외부 시스템과의 관계를 나타냅니다."
)
add_table(
    headers=["Actor / 외부 시스템", "역할 / 데이터 흐름", "방향"],
    rows=[
        ["모바일 사용자 (보행자·자전거)",
         "터치 제스처 입력, 앱 조작, 실감형 3D 지도 열람",
         "User → App"],
        ["GPS 하드웨어 (Android/iOS)",
         "실시간 WGS84 위경도·고도 데이터 제공 (Input.location)",
         "HW → GPSLocationService"],
        ["Cesium Ion 서버",
         "3D Tiles(지형·건물) 스트리밍 — 토큰 인증 후 HTTP 요청",
         "App → Cesium Ion"],
        ["공공 좌표 데이터 (CSV/JSON)",
         "가로등·벤치 등 시설물 위경도 파일, 앱 번들 또는 런타임 로드",
         "File → RoadAssetPlacer"],
        ["자전거 도로 GIS 데이터 (LineString)",
         "자전거 전용 도로 경로 좌표 배열, 런타임 파싱",
         "File → BikeRouteRenderer(예정)"],
    ],
    col_widths=[5, 8, 3],
)
doc.add_paragraph()

add_heading("2.2 Logical Architecture (레이어 구조)", level=2)
doc.add_paragraph(
    "본 프로젝트는 Unity MonoBehaviour 기반 4-레이어 구조로 설계합니다."
)
add_table(
    headers=["레이어", "구성 요소", "책임"],
    rows=[
        ["Presentation\n(View / UI)",
         "MinimapController\nFirstPersonGPSController (OnGUI)\nLocationPermissionHandler (UI Panel)",
         "화면 렌더링, 사용자 입력 수신, UI 상태 표시"],
        ["Application\n(Controller / Service)",
         "FirstPersonGPSController\nGPSLocationService\nRoadAssetPlacer\nBikeRouteRenderer (예정)\nNavigationGuide (예정)",
         "비즈니스 로직, 좌표 변환, 카메라 제어, 데이터 파싱 및 3D 오브젝트 배치"],
        ["Infrastructure\n(외부 연동)",
         "CesiumGeoreference · Cesium3DTileset\nInput.location (GPS)\nUnity AI Navigation (NavMesh)\nPhysics.Raycast",
         "외부 시스템(Cesium Ion, GPS 하드웨어) 접근, 물리 연산"],
        ["Domain\n(Model / Data)",
         "WGS84 좌표 구조체\nCSV/JSON 파싱 결과 DTO\nNavMeshPath / Corners\nRenderTexture (미니맵)",
         "핵심 데이터 모델, 변환 결과 보관"],
    ],
    col_widths=[3.5, 7, 5.5],
)
doc.add_paragraph()

add_heading("2.3 Assembly Definition 구조", level=2)
add_table(
    headers=["Assembly", "포함 스크립트", "외부 의존성"],
    rows=[
        ["Rugem.RoadTools.Runtime",
         "GPSLocationService\nFirstPersonGPSController\nLocationPermissionHandler\nMinimapController\nRoadAssetPlacer",
         "CesiumForUnity\nUnity.AI.Navigation\ncom.unity.inputsystem\nUnity.Mathematics"],
        ["Rugem.RoadTools.Editor",
         "RoadAssetPlacerEditor\niOSBuildPostProcessor",
         "Rugem.RoadTools.Runtime\nUnityEditor"],
        ["(예정) Rugem.BikeRoute.Runtime",
         "BikeRouteRenderer\nBikeRouteShaderController",
         "Rugem.RoadTools.Runtime\nCesiumForUnity"],
        ["(예정) Rugem.Navigation.Runtime",
         "NavigationGuide\nTurnByTurnIndicator",
         "Rugem.RoadTools.Runtime"],
    ],
    col_widths=[4.5, 6.5, 5],
)
doc.add_paragraph()

# ════════════════════════════════════════════════════════════════════════════
# 3. Core Flow (핵심 업무 흐름)
# ════════════════════════════════════════════════════════════════════════════
add_heading("3. Core Flow (핵심 업무 흐름)", level=1)

add_heading("3.1 Flow 1 – GPS 실시간 동기화 & 1인칭 카메라 추적 (US-03, 04, 05)", level=2)
steps_flow1 = [
    "1. 앱 시작",
    "2. LocationPermissionHandler.Start()",
    "   → Android: Permission.RequestUserPermission(FineLocation)",
    "   → iOS: Input.location.Start() 후 상태 폴링",
    "3. 권한 허용 시 OnPermissionGranted 이벤트 발생",
    "4. FirstPersonGPSController.StartGPSTracking()",
    "   → GPSLocationService.StartGPS() 코루틴 시작",
    "5. [Loop: 1초 간격] Input.location.lastData 수신",
    "6. GPSLocationService.ProcessLocationData(lat, lon, alt)",
    "   → WGS84 → ECEF: CesiumWgs84Ellipsoid.LongitudeLatitudeHeightToEarthCenteredEarthFixed()",
    "   → ECEF → Unity World: CesiumGeoreference.TransformECEFToUnity()",
    "   → TargetUnityPosition 갱신, OnRawPositionUpdated 이벤트 발생",
    "7. GPSLocationService.Update(): SmoothedUnityPosition = Lerp(current, target, dt * lerpSpeed)",
    "8. FirstPersonGPSController.OnGPSPositionUpdated()",
    "   → 지면 높이 샘플링: Cesium3DTileset.SampleHeightMostDetailed() (비동기)",
    "   → 실패 시 Physics.Raycast 폴백",
    "   → _targetPosition = (x, groundY + eyeHeight, z)",
    "9. FirstPersonGPSController.Update()",
    "   → transform.position = Lerp(current, _targetPosition, dt * posLerpSpeed)",
    "   → 건물 충돌 감지(0.2초 간격) → 도로 위치로 XZ 고정",
    "10. MinimapController.LateUpdate()",
    "    → 미니맵 카메라가 플레이어 위 400m 높이에서 추적",
]
for step in steps_flow1:
    p = doc.add_paragraph(step)
    p.paragraph_format.left_indent = Cm(0.5)

doc.add_paragraph()

add_heading("3.2 Flow 2 – 커스텀 시설물 자동 배치 (US-06, 07)", level=2)
steps_flow2 = [
    "1. CSV/JSON 파일 파싱 → 시작점(lat1,lon1) / 종료점(lat2,lon2) 좌표 쌍 추출",
    "2. RoadAssetPlacer.BuildNavMesh()",
    "   → Cesium 타일 로드 완료 후 NavMeshSurface.BuildNavMesh() 호출",
    "3. RoadAssetPlacer.PlaceTreeLine(lat1, lon1, lat2, lon2)",
    "   → WGS84 → ECEF → Unity World 1차 변환",
    "   → NavMesh.SamplePosition()으로 노면 최근접점 Snap",
    "   → NavMesh.CalculatePath()로 실제 도로 곡선 경로(Corners) 역추적",
    "4. 각 Corner 사이를 treeInterval(기본 10m) 간격으로 보간",
    "   → Physics.Raycast(down) → 지면 법선(Normal) 획득",
    "   → Prefab Instantiate → position = hit.point",
    "   → CesiumGlobeAnchor 부착 (detectTransformChanges = false)",
    "5. StaticBatchingUtility.Combine() → 드로우 콜 최적화",
    "",
    "   (별도) PlacePointAsset(lat, lon) – 단일 좌표 시설물 배치",
    "   → WGS84 → ECEF → Unity World → Raycast → 지면 안착",
]
for step in steps_flow2:
    p = doc.add_paragraph(step)
    p.paragraph_format.left_indent = Cm(0.5)

doc.add_paragraph()

add_heading("3.3 Flow 3 – 터치 조작 & 모드 전환 (US-11, 12, 13) [예정]", level=2)
steps_flow3 = [
    "1. FirstPersonGPSController.Update() → GetDragDelta()",
    "   → Touchscreen.current.primaryTouch.delta (모바일)",
    "   → Mouse.current.delta (에디터 폴백)",
    "2. 회전 모드 순환: Gyro → Locked → Drag (GUI 버튼 클릭)",
    "   Gyro:  AttitudeSensor → GyroToWorldRotation() → Slerp",
    "   Locked: 현재 rotation 고정",
    "   Drag:  dragYaw/Pitch 누적 → Euler 회전",
    "3. [US-12 예정] 터치 입력 감지 → GPS 카메라 동기화 일시 중단",
    "   → GPSLocationService.StopGPS() 또는 동기화 플래그 false",
    "4. [US-13 예정] '내 위치' 버튼 클릭",
    "   → _targetPosition = GPSLocationService.SmoothedUnityPosition",
    "   → GPS 동기화 재개",
]
for step in steps_flow3:
    p = doc.add_paragraph(step)
    p.paragraph_format.left_indent = Cm(0.5)

doc.add_paragraph()

# ════════════════════════════════════════════════════════════════════════════
# 4. 데이터 설계
# ════════════════════════════════════════════════════════════════════════════
add_heading("4. 데이터 설계", level=1)

add_heading("4.1 주요 데이터 흐름 (Data Flow)", level=2)
add_table(
    headers=["데이터", "형식", "출처", "소비자"],
    rows=[
        ["GPS 위경도·고도", "double lat/lon/alt", "Input.location (OS)", "GPSLocationService"],
        ["WGS84 → ECEF", "double3 (ECEF)", "CesiumWgs84Ellipsoid", "GPSLocationService"],
        ["ECEF → Unity World", "Vector3 / double3", "CesiumGeoreference", "FirstPersonGPSController, RoadAssetPlacer"],
        ["Cesium 지형 고도", "double (height)", "Cesium3DTileset.SampleHeightMostDetailed()", "FirstPersonGPSController"],
        ["시설물 좌표 데이터", "CSV / JSON (lat, lon, type)", "외부 공공 데이터 파일", "RoadAssetPlacer"],
        ["NavMesh 경로", "NavMeshPath.corners[]", "NavMesh.CalculatePath()", "RoadAssetPlacer"],
        ["자전거 도로 경로", "LineString (double[] pairs)", "GIS 공공 데이터 (예정)", "BikeRouteRenderer (예정)"],
        ["미니맵 렌더", "RenderTexture (256×256)", "MinimapCamera", "MinimapController.OnGUI()"],
    ],
    col_widths=[4, 3.5, 5, 4.5],
)
doc.add_paragraph()

add_heading("4.2 주요 컴포넌트 관계", level=2)
doc.add_paragraph(
    "아래는 런타임 시 주요 MonoBehaviour 간의 참조 관계를 나타냅니다."
)
add_table(
    headers=["컴포넌트", "참조 대상", "참조 방식"],
    rows=[
        ["FirstPersonGPSController", "GPSLocationService", "SerializeField + 이벤트 구독 (OnRawPositionUpdated)"],
        ["FirstPersonGPSController", "LocationPermissionHandler", "SerializeField + 이벤트 구독"],
        ["FirstPersonGPSController", "Cesium3DTileset (WorldTerrain)", "SerializeField (async 높이 샘플링)"],
        ["FirstPersonGPSController", "CesiumGlobeAnchor", "AddComponent (자동 부착)"],
        ["GPSLocationService", "CesiumGeoreference", "SerializeField / FindAnyObjectByType"],
        ["RoadAssetPlacer", "CesiumGeoreference", "GetComponentInParent / FindAnyObjectByType"],
        ["RoadAssetPlacer", "NavMeshSurface", "GetComponent / AddComponent"],
        ["MinimapController", "FirstPersonGPSController", "FindAnyObjectByType (followTarget)"],
        ["MinimapController", "CesiumCameraManager", "CesiumCameraManager.GetOrCreate()"],
    ],
    col_widths=[5, 4.5, 7],
)
doc.add_paragraph()

add_heading("4.3 Integrity Rules (무결성 규칙)", level=2)
rules = [
    "CesiumGeoreference는 씬에 반드시 1개만 존재해야 하며, 모든 좌표 변환의 단일 기준점이다.",
    "CesiumGlobeAnchor.detectTransformChanges = false — 런타임 Transform 변경 시 Cesium 내부 좌표 재계산을 막아 성능 보호.",
    "NavMesh는 Cesium 타일 로드 완료 이후에만 BuildNavMesh()를 호출해야 한다 (정적 베이크 불가).",
    "GPS 고도는 노이즈가 심하므로 카메라 Y축은 반드시 Cesium SampleHeight 또는 Raycast 결과를 사용한다.",
    "StaticBatching은 배치 완료 후 1회만 적용한다 (런타임 추가 오브젝트에 재적용 금지).",
]
for r in rules:
    p = doc.add_paragraph(f"• {r}")

doc.add_paragraph()

# ════════════════════════════════════════════════════════════════════════════
# 5. 앞으로 개발해야 할 요소 (Sprint 2·3 상세 설계)
# ════════════════════════════════════════════════════════════════════════════
add_heading("5. 앞으로 개발해야 할 요소 (Sprint 2 · 3)", level=1)

add_heading("5.1 자전거 도로 렌더링 시스템 (US-09, US-10) — Sprint 2", level=2)
add_table(
    headers=["항목", "내용"],
    rows=[
        ["대상 클래스 (신규)", "BikeRouteRenderer.cs\n(Rugem.BikeRoute.Runtime Assembly 분리 예정)"],
        ["입력 데이터", "LineString JSON: [[lon1,lat1],[lon2,lat2],…] 형식 공공 자전거 도로 데이터"],
        ["좌표 변환",
         "각 좌표 쌍 → CesiumWgs84Ellipsoid.LLH→ECEF → CesiumGeoreference.ECEF→Unity World"],
        ["메쉬 생성 (US-09)",
         "Vector3[] 배열로 변환된 World 좌표를 바탕으로 MeshFilter/MeshRenderer 폴리라인 메쉬 생성\n"
         "도로 폭(width) Inspector 조절 가능, 지면 밀착(Raycast Y 보정) 적용"],
        ["셰이더 (US-10)",
         "URP Shader Graph 또는 HLSL Custom Pass:\n"
         "  • 발광(Emission) 효과: HDR Color 설정으로 야간 식별성 확보\n"
         "  • 점선 애니메이션: _Time 기반 UV 스크롤로 진행 방향 시각화\n"
         "  • LOD: 카메라 거리 50m 이하에서만 발광 활성화"],
        ["CesiumGlobeAnchor", "메쉬 부모 오브젝트에 부착, detectTransformChanges = false"],
    ],
    col_widths=[4, 12],
)
doc.add_paragraph()

add_heading("5.2 경로 안내 시각화 시스템 (US-14) — Sprint 2", level=2)
add_table(
    headers=["항목", "내용"],
    rows=[
        ["대상 클래스 (신규)", "NavigationGuide.cs, TurnByTurnIndicator.cs\n(Rugem.Navigation.Runtime 예정)"],
        ["경로 계산",
         "목적지 WGS84 좌표 입력 → NavMesh.CalculatePath() 또는 외부 라우팅 API(선택)\n"
         "결과 corners[] → Unity World 좌표 배열로 변환"],
        ["바닥 인디케이터",
         "1인칭 시점 바닥면에 DecalProjector (URP Decal) 또는 투명 Plane 메쉬로\n"
         "화살표 방향 표시 (플레이어 전방 5~15m 구간에만 표시)"],
        ["Turn-by-Turn",
         "방향 전환 지점(corner) 10m 이내 접근 시 OnGUI 또는 Canvas WorldSpace로 방향 아이콘 표출\n"
         "직진 / 좌회전 / 우회전 아이콘 3종"],
        ["GPS 연동",
         "GPSLocationService.SmoothedUnityPosition을 매 프레임 소비하여\n"
         "경로 경과 구간 판정 및 다음 Corner로 인덱스 전진"],
    ],
    col_widths=[4, 12],
)
doc.add_paragraph()

add_heading("5.3 UI/UX 완성 (US-11 심화, US-12, US-13) — Sprint 2 · 3", level=2)
add_table(
    headers=["User Story", "설계 요소", "구현 방안"],
    rows=[
        ["US-11 터치 기반 시점 조작",
         "1손가락 스와이프 → 카메라 회전\n2손가락 핀치 → FOV Zoom",
         "현재: RotationMode.Drag (스와이프 회전) 구현 완료\n추가: Touchscreen.current의 primaryTouch·secondaryTouch Pinch 델타로 Camera.fieldOfView 조절"],
        ["US-12 조작 모드 심리스 전환",
         "터치 감지 즉시 GPS 동기화 일시 중단",
         "FirstPersonGPSController에 _isManualMode 플래그 추가\n터치 입력 시 _isManualMode = true → GPS 위치 적용 스킵\n일정 시간(3초) 터치 없으면 자동 복귀"],
        ["US-13 '내 위치' 복귀",
         "UI 버튼 → GPS 좌표로 즉시 이동",
         "OnGUI 또는 Canvas Button\n클릭 시 _isManualMode = false, _targetPosition = SmoothedUnityPosition\nGPSLocationService 동기화 재개"],
    ],
    col_widths=[3, 5.5, 7.5],
)
doc.add_paragraph()

add_heading("5.4 시설물 가시성 필터링 (US-08) — Sprint 3", level=2)
add_table(
    headers=["항목", "내용"],
    rows=[
        ["설계 위치", "RoadAssetPlacer 또는 별도 AssetLayerManager.cs"],
        ["레이어 구조",
         "가로등 → Unity Layer 'StreetLight'\n벤치 → Unity Layer 'Bench'\n버스정류장 → Unity Layer 'BusStop'"],
        ["On/Off 방법",
         "Camera.cullingMask 비트 조작 또는 GameObject.SetActive()\n"
         "UI: Toggle 그룹 (Canvas 또는 OnGUI)"],
        ["성능 고려",
         "SetActive()는 GC 부하 → 대신 Renderer.enabled 사용\n"
         "레이어별 GameObject 풀링으로 생성/파괴 비용 최소화"],
    ],
    col_widths=[4, 12],
)
doc.add_paragraph()

add_heading("5.5 모바일 성능 최적화 (US-15) — Sprint 3", level=2)
add_table(
    headers=["최적화 기법", "대상", "구현 방안"],
    rows=[
        ["거리 기반 LOD", "커스텀 시설물 프리팹",
         "Camera.main.WorldToViewportPoint()로 거리 계산\n50m 이상: LOD 1 (저폴리), 100m 이상: SetActive(false)"],
        ["Static Batching", "RoadAssetPlacer 배치 완료 객체",
         "StaticBatchingUtility.Combine(lineParent) — 현재 구현 완료"],
        ["GPU Instancing", "동일 프리팹(가로등 등) 다수 배치",
         "Material.enableInstancing = true\nRenderer 공유 Material 사용"],
        ["3D Tiles 스트리밍 제한", "Cesium3DTileset",
         "maximumScreenSpaceError 조정 (기본 16 → 모바일 32)\nCesiumCameraManager에 미니맵 카메라만 추가 등록하여 불필요 타일 로드 억제"],
        ["배터리 효율", "GPS 폴링 간격",
         "이동 속도 < 0.5 m/s 감지 시 _pollIntervalSeconds 3초로 자동 증가\n이동 중 복귀"],
        ["Render Pipeline 설정", "URP Mobile_RPAsset",
         "MSAA 비활성화, Shadow Distance 축소 (100m)\nPost-Processing 최소화"],
    ],
    col_widths=[4, 3.5, 8.5],
)
doc.add_paragraph()

# ════════════════════════════════════════════════════════════════════════════
# 6. Development Environment
# ════════════════════════════════════════════════════════════════════════════
add_heading("6. Development Environment", level=1)
add_table(
    headers=["항목", "값 / 버전", "비고"],
    rows=[
        ["Unity",          "2022.3 LTS",                   "장기 지원 버전"],
        ["C#",             ".NET Standard 2.1",             "Unity 스크립팅 백엔드 IL2CPP"],
        ["Render Pipeline","Universal Render Pipeline (URP)","PC_RPAsset / Mobile_RPAsset 분리 설정"],
        ["Cesium for Unity","최신 패키지 (com.cesium.unity)", "Cesium Ion 토큰 인증 필요"],
        ["Unity Input System","com.unity.inputsystem",      "새 Input System (Touchscreen, AttitudeSensor)"],
        ["Unity AI Navigation","com.unity.ai.navigation",   "런타임 NavMesh 빌드"],
        ["Unity Mathematics","com.unity.mathematics",       "double3 정밀도 연산"],
        ["Android SDK",    "API Level 24 (Android 7.0) 이상","AndroidManifest.xml 위치 권한 선언"],
        ["iOS",            "iOS 13.0 이상",                 "Xcode 빌드, NSLocationWhenInUse 권한"],
        ["개발 OS",        "Windows 11 / macOS",            "Windows에서 iOS 빌드는 Mac 필요"],
        ["버전 관리",      "Git (GitHub)",                  "브랜치 전략: main / app_beta1"],
        ["IDE",            "JetBrains Rider / VS 2022",     "Assembly Definition 기반 프로젝트 구성"],
        ["빌드 자동화",    "Unity Build Settings",          "Android: APK / iOS: Xcode Project"],
    ],
    col_widths=[4, 5, 7],
)
doc.add_paragraph()

add_heading("6.1 프로젝트 폴더 구조", level=2)
folder_lines = [
    "Assets/",
    "├── CesiumSettings/          # Cesium Ion 서버·런타임 설정 에셋",
    "├── Darth_Artisan/Free_Trees/ # 가로수 프리팹 (Fir, Oak, Palm, Poplar)",
    "├── Plugins/Android/         # AndroidManifest.xml",
    "├── RoadTools/",
    "│   ├── Editor/              # RoadAssetPlacerEditor.cs, iOSBuildPostProcessor.cs",
    "│   └── Runtime/",
    "│       ├── GPS/             # GPSLocationService, FirstPersonGPSController,",
    "│       │                    #   LocationPermissionHandler, MinimapController",
    "│       └── RoadAssetPlacer.cs",
    "├── Scenes/",
    "│   ├── SampleScene.unity    # 기본 씬",
    "│   ├── level1.unity         # 스프린트 1 테스트 씬",
    "│   └── level2.unity",
    "├── Settings/                # URP RPAsset (PC/Mobile), Renderer 설정",
    "├── NavMesh-dorohe.asset     # 사전 베이크 NavMesh (Fallback)",
    "└── InputSystem_Actions.inputactions  # 새 Input System 액션 맵",
    "",
    "(예정) Assets/BikeRoute/     # BikeRouteRenderer, 셰이더",
    "(예정) Assets/Navigation/    # NavigationGuide, TurnByTurnIndicator",
]
for line in folder_lines:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    run = p.add_run(line)
    run.font.name = 'Courier New'
    run.font.size = Pt(9)

doc.add_paragraph()

# ════════════════════════════════════════════════════════════════════════════
# 7. Traceability
# ════════════════════════════════════════════════════════════════════════════
add_heading("7. Traceability (요구사항 추적)", level=1)
add_table(
    headers=["US", "User Story", "Epic", "설계 요소 (클래스/메서드)", "Sprint", "상태"],
    rows=[
        ["US-01", "Cesium Ion 서버 연동",        "E1", "CesiumGeoreference, Cesium3DTileset 씬 설정",                     "1", "완료"],
        ["US-02", "대상 지역 데이터 최적화",      "E1", "Cesium3DTileset 타일 우선순위, 고도 이격 수정",                   "1", "완료"],
        ["US-03", "모바일 위치 권한 획득",        "E2", "LocationPermissionHandler.CheckAndRequestPermission()",          "1", "완료"],
        ["US-04", "지리 좌표계 변환 엔진",        "E2", "GPSLocationService.ConvertToUnityPosition(), ProcessLocationData()", "1", "완료"],
        ["US-05", "1인칭 카메라 위치 동기화",    "E2", "FirstPersonGPSController.OnGPSPositionUpdated(), SampleAndUpdateGroundHeight()", "1", "완료"],
        ["US-06", "외부 좌표 데이터 파싱",        "E3", "RoadAssetPlacer (CSV 파싱 연동)",                                "1", "완료"],
        ["US-07", "데이터 기반 객체 자동 배치",  "E3", "RoadAssetPlacer.PlaceTreeLine(), PlacePointAsset()",             "1", "완료"],
        ["US-08", "시설물 가시성 필터링",         "E3", "AssetLayerManager (예정), Renderer.enabled",                    "3", "예정"],
        ["US-09", "자전거 도로 폴리라인 생성",   "E4", "BikeRouteRenderer.GenerateMesh() (예정)",                        "2", "예정"],
        ["US-10", "자전거 도로 특화 셰이더",     "E4", "BikeRouteShaderController, URP Shader Graph (예정)",             "2", "예정"],
        ["US-11", "터치 기반 시점 조작",          "E5", "FirstPersonGPSController.GetDragDelta(), CycleRotationMode()",   "2", "부분완료"],
        ["US-12", "조작 모드 심리스 전환",        "E5", "FirstPersonGPSController._isManualMode (예정)",                  "3", "예정"],
        ["US-13", "'내 위치' 복귀 기능",          "E5", "OnGUI 버튼 → _targetPosition 초기화 (예정)",                    "3", "예정"],
        ["US-14", "실시간 경로 가이드 시각화",   "E6", "NavigationGuide, TurnByTurnIndicator (예정)",                    "2", "예정"],
        ["US-15", "거리 기반 LOD 및 컬링",        "E7", "LOD 컴포넌트, Camera.cullingMask, Mobile_RPAsset 설정 (예정)",   "3", "예정"],
    ],
    col_widths=[1.5, 4.5, 1.5, 6.5, 1.5, 1.5],
)
doc.add_paragraph()

# 저장
doc.save(OUTPUT)
print(f"저장 완료: {OUTPUT}")
