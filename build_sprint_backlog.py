import shutil, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ── 유틸리티 ─────────────────────────────────────────────────────────────────

def shd(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement('w:shd')
    s.set(qn('w:val'), 'clear')
    s.set(qn('w:color'), 'auto')
    s.set(qn('w:fill'), hex_color)
    tcPr.append(s)

def set_col_width(table, col_idx, width_cm):
    for row in table.rows:
        row.cells[col_idx].width = Cm(width_cm)

def header_row(table, headers, bg='4472C4', fg='FFFFFF'):
    row = table.rows[0]
    for i, h in enumerate(headers):
        cell = row.cells[i]
        cell.text = h
        shd(cell, bg)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor.from_string(fg)
                r.font.size = Pt(9)

def data_rows(table, rows, alt='D9E2F3'):
    for idx, row_data in enumerate(rows):
        row = table.add_row()
        for i, val in enumerate(row_data):
            cell = row.cells[i]
            cell.text = str(val)
            if idx % 2 == 0:
                shd(cell, alt)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)

def task_section(doc, heading, tasks):
    try:
        p = doc.add_paragraph(heading, style='new 스타일1')
    except Exception:
        p = doc.add_paragraph(heading)
        for r in p.runs:
            r.bold = True; r.font.size = Pt(12)

    t = doc.add_table(rows=1, cols=5)
    t.style = 'Normal Table'
    header_row(t, ['Task ID', 'Task', 'Owner', 'Est(h)', 'Done Criteria'])
    data_rows(t, tasks)
    # 컬럼 폭 설정
    for w, ci in zip([1.3, 6.5, 1.4, 1.2, 5.2], range(5)):
        set_col_width(t, ci, w)
    doc.add_paragraph()

# ── 템플릿 복사 및 예시 내용 제거 ──────────────────────────────────────────

shutil.copy("[양식 6-2] Sprint Backlog 양식 (1).docx", "Sprint_Backlog_Sprint1_완성.docx")
doc = Document("Sprint_Backlog_Sprint1_완성.docx")

# 커버·개정이력 이후 모든 요소 삭제
body = doc.element.body
rev_table_elem = doc.tables[1]._element
all_elems = list(body)
for elem in all_elems[all_elems.index(rev_table_elem) + 1:]:
    body.remove(elem)

# ── 1. 메인 헤딩 ─────────────────────────────────────────────────────────────

doc.add_heading("Sprint Backlog – 3D 디지털 트윈 1인칭 내비게이션 시스템 (Sprint 1)", level=1)

# ── 1. Sprint Goal ────────────────────────────────────────────────────────────

doc.add_heading("1. Sprint Goal", level=2)
doc.add_paragraph(
    "Cesium 지형 위 GPS 연동 1인칭 카메라 이동, 외부 좌표 기반 시설물 자동 배치, "
    "탑뷰 미니맵이 동작 가능한 모바일 앱 상태를 구축한다.",
    style='First Paragraph'
)

# ── 2. Selected User Stories ──────────────────────────────────────────────────

doc.add_heading("2. Selected User Stories", level=2)
us_tbl = doc.add_table(rows=1, cols=2)
us_tbl.style = 'Normal Table'
header_row(us_tbl, ['ID', 'Title'])
data_rows(us_tbl, [
    ('US-01', 'Cesium ion 서버 연동'),
    ('US-02', '대상 지역 데이터 최적화'),
    ('US-03', '모바일 위치 권한 획득'),
    ('US-04', '지리좌표계 변환 엔진'),
    ('US-05', '1인칭 카메라 위치 동기화'),
    ('US-06', '외부 좌표 데이터 파싱'),
    ('US-07', '데이터 기반 객체 자동 배치'),
    ('US-08', '시설물 가시성 필터링'),
    ('US-11', '터치 기반 시점 조작'),
    ('US-12', '조작 모드 심리스 전환'),
    ('US-16', 'Cesium 지형 표면 높이 정밀 측정'),
    ('US-17', 'GPS 고도 노이즈 차단 및 지면 높이 캐싱'),
    ('US-18', '건물 내부 침투 감지 및 인접 도로 위치 고정'),
    ('US-19', '탑뷰 미니맵 (Cesium 타일 연동)'),
    ('US-20', '앱 백그라운드·포커스 복귀 시 센서 자동 재활성화'),
])
set_col_width(us_tbl, 0, 2.5); set_col_width(us_tbl, 1, 10.0)

# ── 3. Task Breakdown ─────────────────────────────────────────────────────────

doc.add_heading("3. Task Breakdown (Detailed)", level=2)

task_section(doc, "US-01  Cesium ion 서버 연동", [
    ('T01', 'Cesium for Unity v1.23.1 패키지 설치 및 Ion 액세스 토큰 등록', 'PO', 2, 'Ion 토큰 인증 성공, 콘솔 오류 없음'),
    ('T02', 'CesiumGeoreference 오브젝트 배치 및 기준 좌표(위경도) 설정', 'PO', 2, '대상 지역 중심 좌표 원점 설정 확인'),
    ('T03', 'Cesium World Terrain 및 OSM 건물 Tileset 로드 확인', 'PO', 2, '지형·건물 3D 타일 에디터 렌더링 확인'),
    ('T04', 'Assembly Definition(asmdef) 구성 및 Cesium·InputSystem GUID 참조 등록', 'PO', 2, '컴파일 에러 0건'),
])

task_section(doc, "US-02  대상 지역 데이터 최적화", [
    ('T05', '타겟 지역 중심 좌표 기반 타일 우선 로딩 설정', 'PO', 3, '대상 지역 타일 우선 스트리밍 확인'),
    ('T06', '지형-건물 고도 이격 현상 원인 분석 및 오프셋 수정', 'PO', 4, '지형·건물 고도 일치 육안 확인'),
])

task_section(doc, "US-03  모바일 위치 권한 획득", [
    ('T07', 'LocationPermissionHandler.cs 작성 — Android Permission.FineLocation 요청 콜백 구현', 'SM', 3, '실기기 권한 팝업 정상 표시'),
    ('T08', 'iOS LocationService.Start() 기반 권한 확인 코루틴 구현', 'SM', 2, 'iOS 권한 획득 흐름 확인'),
    ('T09', '권한 거부 시 안내 패널 및 앱 설정 이동 버튼 UI 구현', 'Dev1', 2, '거부 시 패널 표시, 설정 화면 이동 확인'),
])

task_section(doc, "US-04  지리좌표계 변환 엔진", [
    ('T10', 'GPSLocationService.cs 작성 — WGS84 → ECEF → Unity 3단계 변환 로직 구현', 'SM', 4, '변환 좌표 logcat 출력 확인'),
    ('T11', 'Vector3.Lerp 위치 보간 적용으로 GPS 업데이트 시 카메라 떨림 방지', 'SM', 3, '이동 중 부드러운 보간 육안 확인'),
    ('T12', 'GPS 폴링 루프(1초 간격) 및 서비스 중단 시 자동 재시작 구현', 'SM', 3, '서비스 중단·재시작 로그 확인'),
    ('T13', '에디터 GPS 시뮬레이션 모드 구현 (CesiumGeoreference 원점 좌표 사용)', 'SM', 2, '에디터 실행 시 시뮬레이션 위치 로그 확인'),
])

task_section(doc, "US-05  1인칭 카메라 위치 동기화", [
    ('T14', 'FirstPersonGPSController.cs 기본 구조 작성 — GPS 이벤트 수신 및 CesiumGlobeAnchor 연동', 'SM', 3, 'GPS 위치 수신 시 카메라 이동 확인'),
    ('T15', 'CesiumGlobeAnchor 위치 Lerp 보간 구현', 'SM', 3, '카메라 부드러운 이동 확인'),
    ('T16', 'Awake()에서 CesiumCameraController 강제 비활성화 (위치·회전 충돌 방지)', 'SM', 2, '위치 충돌 현상 미발생 및 로그 스팸 제거 확인'),
    ('T17', 'InputSystem.EnableDevice(AttitudeSensor.current) 자이로 센서 초기화', 'SM', 2, 'AttitudeSensor 활성화 로그 확인'),
])

task_section(doc, "US-06  외부 좌표 데이터 파싱", [
    ('T18', 'CSV 파일 읽기 및 위경도·속성 컬럼 파싱 모듈 작성', 'PO', 3, '파싱된 좌표 데이터 로그 출력 확인'),
    ('T19', '파일 미존재·형식 오류 예외 처리 및 Debug.LogError 출력 구현', 'PO', 2, '잘못된 파일 입력 시 오류 로그 출력 확인'),
])

task_section(doc, "US-07  데이터 기반 객체 자동 배치", [
    ('T20', 'RoadAssetPlacer.cs 기본 구조 작성 — 좌표 기반 프리팹 인스턴싱', 'PO', 3, '지정 좌표에 시설물 프리팹 생성 확인'),
    ('T21', '런타임 NavMeshSurface.BuildNavMesh() 동적 호출 구현', 'PO', 4, '런타임 NavMesh 빌드 성공 로그 확인'),
    ('T22', 'NavMesh.CalculatePath Corners 기반 곡선 경로 보간 배치 포인트 생성', 'PO', 4, '직선 좌표 간 곡선 경로 Corner 데이터 확보 확인'),
    ('T23', '수직 하방 Raycast 법선 벡터 산출 및 객체 경사 정렬 적용', 'PO', 3, '경사 지형 위 시설물 기울기 정렬 육안 확인'),
    ('T24', 'CesiumGlobeAnchor 부착 및 detectTransformChanges=false 최적화', 'PO', 2, '대량 배치 시 Transform 감지 비용 제거 확인'),
])

task_section(doc, "US-08  시설물 가시성 필터링", [
    ('T25', '시설물 타입별 레이어 분류 및 SetActive On/Off 토글 기능 구현', 'Dev1', 3, '타입 선택 시 해당 시설물만 표시·숨김 확인'),
])

task_section(doc, "US-11  터치 기반 시점 조작", [
    ('T26', 'Touchscreen.primaryTouch.press.isPressed 기반 1손가락 드래그 회전 구현', 'SM', 3, '실기기 1손가락 스와이프로 카메라 회전 확인'),
    ('T27', '에디터 마우스 좌클릭 드래그 폴백 구현', 'SM', 2, '에디터 마우스 드래그 회전 동작 확인'),
])

task_section(doc, "US-12  조작 모드 심리스 전환", [
    ('T28', 'RotationMode enum(Gyro/Locked/Drag) 및 CycleRotationMode() 구현', 'SM', 2, '버튼 클릭 시 3모드 순환 전환 확인'),
    ('T29', '화면 비례형 모드 토글 버튼 OnGUI() 구현 (우측 상단, 미니맵 하단 배치)', 'Dev1', 2, '모드 표시 버튼 렌더링 및 탭 반응 확인'),
    ('T30', 'Locked 모드 회전값 고정, Drag 모드 진입 시 Yaw·Pitch 초기화 구현', 'SM', 2, '각 모드 전환 시 회전 상태 정확히 유지 확인'),
])

task_section(doc, "US-16  Cesium 지형 표면 높이 정밀 측정", [
    ('T31', 'SampleAndUpdateGroundHeight() async void 메서드 작성', 'SM', 4, '비동기 호출 완료 후 지면 높이 갱신 로그 확인'),
    ('T32', 'Cesium3DTileset.SampleHeightMostDetailed() API 연동 및 타일 고도→Unity 좌표 변환', 'SM', 3, '샘플링 성공 시 고도값·Unity Y 좌표 로그 확인'),
    ('T33', 'Cesium 샘플링 실패 시 Physics.Raycast 폴백 및 200m 이상 이상값 필터링', 'SM', 3, '폴백 전환 로그 및 이상값 거부 로그 확인'),
    ('T34', '_samplingHeight 플래그로 비동기 중복 호출 방지 구현', 'SM', 2, '동시 다중 호출 없음 확인'),
])

task_section(doc, "US-17  GPS 고도 노이즈 차단 및 지면 높이 캐싱", [
    ('T35', '_cachedGroundY 캐싱 변수 및 초기값(float.MinValue) 설정', 'SM', 2, '최초 GPS 수신 시 캐시 저장 확인'),
    ('T36', '수평 2m 이상 이동 시에만 지면 재감지 임계값(GroundCheckMoveThreshold) 적용', 'SM', 2, '소폭 이동 시 캐시 유지, 2m 이상 이동 시 재감지 로그 확인'),
])

task_section(doc, "US-18  건물 내부 침투 감지 및 인접 도로 위치 고정", [
    ('T37', 'IsInsideBuilding() 수평 4방향 레이(15m) 전차단 판정 구현', 'SM', 3, '건물 내부 진입 시 감지 로그 출력 확인'),
    ('T38', 'FindNearestRoadXZ() 방사형 탐색(2m 간격, 최대 30m) 구현', 'SM', 4, '인접 개방 위치 좌표 반환 로그 확인'),
    ('T39', 'HandleBuildingCollision() XZ 고정 및 충돌 해제 시 잠금 해제 로직 구현', 'SM', 2, '건물 탈출 시 자유 이동 복귀 확인'),
    ('T40', '0.2초 간격 체크 타이머 적용으로 매 프레임 레이캐스트 비용 절감', 'SM', 2, '프레임 드롭 없이 충돌 감지 동작 확인'),
])

task_section(doc, "US-19  탑뷰 미니맵 (Cesium 타일 연동)", [
    ('T41', 'MinimapController.cs 작성 — 직교 카메라 및 RenderTexture 생성', 'Dev1', 3, '미니맵 영역 렌더링 확인'),
    ('T42', 'CesiumCameraManager.additionalCameras에 미니맵 카메라 등록', 'Dev1', 2, '미니맵 frustum 내 Cesium 타일 전체 스트리밍 확인'),
    ('T43', 'OnGUI() 미니맵 RenderTexture 우측 상단 표시 및 테두리 렌더링', 'Dev1', 2, '화면 우측 상단 미니맵 UI 표시 확인'),
    ('T44', '플레이어 방향 화살표 삼각형 텍스처 생성 및 Yaw 기반 회전 렌더링', 'Dev1', 2, '이동 방향에 따른 화살표 회전 확인'),
    ('T45', 'LateUpdate() 플레이어 위치 추적 및 미니맵 카메라 XZ 동기화', 'Dev1', 2, '플레이어 이동 시 미니맵 중심 추적 확인'),
])

task_section(doc, "US-20  앱 백그라운드·포커스 복귀 시 센서 자동 재활성화", [
    ('T46', 'OnApplicationPause() 백그라운드 전환 시 AttitudeSensor 비활성화 구현', 'SM', 2, '백그라운드 전환 시 센서 비활성화 로그 확인'),
    ('T47', 'OnApplicationFocus() 포커스 복귀 시 AttitudeSensor 재활성화 구현', 'SM', 1, '앱 복귀 시 자이로 회전 즉시 재개 확인'),
])

# ── 4. Definition of Done Mapping ─────────────────────────────────────────────

doc.add_heading("4. Definition of Done Mapping", level=2)
dod_tbl = doc.add_table(rows=1, cols=4)
dod_tbl.style = 'Normal Table'
header_row(dod_tbl, ['User Story', 'AC', '검증 Task', '검증 결과'])
data_rows(dod_tbl, [
    ('US-01', 'AC-01  Ion 토큰 인증 성공',                           'T01 (토큰 등록)',          'Ion 콘솔 인증 토큰 유효 확인'),
    ('US-01', 'AC-02  World Terrain·OSM 건물 로드 확인',              'T03 (타일셋 로드)',        '지형·건물 3D 타일 렌더링 확인'),
    ('US-02', 'AC-01  대상 지역 타일 우선 로딩 설정',                  'T05 (우선 로딩)',          '대상 지역 타일 우선 스트리밍 확인'),
    ('US-02', 'AC-02  지형-건물 고도 이격 현상 수정',                  'T06 (고도 수정)',          '고도 일치 육안 확인'),
    ('US-03', 'AC-01  Android/iOS 권한 요청 팝업 구현',               'T07, T08 (권한 요청)',    '실기기 권한 팝업 표시 확인'),
    ('US-03', 'AC-02  권한 거부 시 설정 안내 UI',                     'T09 (안내 UI)',            '거부 시 패널·설정 이동 버튼 동작 확인'),
    ('US-04', 'AC-01  WGS84 → Unity 좌표 변환 성공',                  'T10 (변환 로직)',          '변환 좌표값 logcat 출력 확인'),
    ('US-04', 'AC-02  카메라 떨림 방지 Lerp 적용',                    'T11 (Lerp 보간)',          '이동 중 부드러운 보간 육안 확인'),
    ('US-05', 'AC-01  GPS 업데이트 시 카메라 Transform 동기화',        'T14, T15 (카메라 연동)',  'GPS 이동 시 카메라 위치 반영 확인'),
    ('US-05', 'AC-02  CesiumCameraController 충돌 제거',              'T16 (충돌 방지)',          '위치 충돌 현상 미발생 확인'),
    ('US-06', 'AC-01  위경도·속성 데이터 파싱 성공',                   'T18 (파싱 모듈)',          '파싱 데이터 로그 출력 확인'),
    ('US-06', 'AC-02  파싱 실패 예외 처리',                           'T19 (예외 처리)',          '오류 파일 입력 시 LogError 출력 확인'),
    ('US-07', 'AC-01  프리팹이 지정 좌표에 인스턴싱됨',               'T20 (인스턴싱)',           '지정 좌표 시설물 배치 육안 확인'),
    ('US-07', 'AC-02  지형 경사도 기반 객체 회전 정렬',               'T23 (법선 정렬)',          '경사면 시설물 기울기 정렬 육안 확인'),
    ('US-08', 'AC-01  시설물 타입별 레이어 On/Off',                   'T25 (필터링)',             '타입 선택 시 가시성 토글 확인'),
    ('US-11', 'AC-01  1손가락 스와이프 카메라 회전',                   'T26 (터치 드래그)',        '실기기 드래그 회전 동작 확인'),
    ('US-12', 'AC-01  3모드(Gyro/Locked/Drag) 순환 전환',            'T28, T29 (모드 전환)',    '버튼 탭 시 모드 전환 확인'),
    ('US-12', 'AC-02  모드별 회전 상태 정확히 유지',                   'T30 (상태 관리)',          '각 모드 회전값 정확히 유지 확인'),
    ('US-16', 'AC-01  Cesium 타일 고도 샘플링 성공',                  'T31, T32 (높이 샘플링)', '샘플링 성공 고도값 로그 확인'),
    ('US-16', 'AC-02  샘플링 실패 시 Raycast 폴백 동작',              'T33 (폴백)',               '폴백 전환 로그 확인'),
    ('US-17', 'AC-01  지면 높이 캐싱으로 Y축 떨림 방지',              'T35, T36 (캐싱)',          '소폭 이동 시 Y축 고정 육안 확인'),
    ('US-18', 'AC-01  건물 내부 진입 감지',                           'T37 (레이 판정)',          '건물 진입 시 감지 로그 출력 확인'),
    ('US-18', 'AC-02  인접 도로 위치로 XZ 고정',                      'T38, T39 (XZ 고정)',      '건물 내부 XZ 고정·탈출 시 해제 확인'),
    ('US-19', 'AC-01  미니맵 전 영역 Cesium 타일 렌더링',             'T42 (카메라 등록)',        '미니맵 전 범위 지형 표시 확인'),
    ('US-19', 'AC-02  플레이어 방향 화살표 표시',                     'T44 (화살표)',             '이동 방향 화살표 회전 육안 확인'),
    ('US-20', 'AC-01  백그라운드·포커스 복귀 시 자이로 재개',         'T46, T47 (센서 재활성)', '앱 복귀 후 자이로 회전 정상 재개 확인'),
])
for w, ci in zip([2.2, 6.0, 3.5, 5.0], range(4)):
    set_col_width(dod_tbl, ci, w)

doc.save("Sprint_Backlog_Sprint1_완성.docx")
print("✅ Sprint_Backlog_Sprint1_완성.docx 생성 완료")
